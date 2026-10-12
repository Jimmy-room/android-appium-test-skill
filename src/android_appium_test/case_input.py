"""提取用例文件的可读内容，语义转换由 Codex Agent 完成。"""

from __future__ import annotations

import json
from pathlib import Path

from .config import ConfigError


def read_case_input(path: str | Path) -> str:
    """读取 Markdown、文本、Word、Excel 或结构化文件的原文。"""
    source = Path(path)
    if not source.is_file():
        raise ConfigError(f"用例文件不存在：{source}")
    suffix = source.suffix.lower()
    if suffix in {".txt", ".md", ".yaml", ".yml", ".json"}:
        content = source.read_text(encoding="utf-8-sig")
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ConfigError("读取 Word 需要 python-docx；也可提供 Markdown") from exc
        document = Document(source)
        rows = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            rows.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        content = "\n".join(rows)
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ConfigError("读取 Excel 需要 openpyxl；也可提供 Markdown") from exc
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            content = "\n".join(
                f"[{sheet.title}] " + " | ".join("" if cell is None else str(cell) for cell in row)
                for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)
            )
        finally:
            workbook.close()
    else:
        raise ConfigError(f"暂不支持用例格式：{suffix}")
    if not content.strip():
        raise ConfigError("用例文件为空")
    return content


def is_structured_spec(path: str | Path) -> bool:
    """只识别明确含 features 的规范文件，不猜测自然语言含义。"""
    source = Path(path)
    if source.suffix.lower() not in {".json", ".yaml", ".yml"}:
        return False
    try:
        if source.suffix.lower() == ".json":
            data = json.loads(source.read_text(encoding="utf-8-sig"))
        else:
            import yaml
            data = yaml.safe_load(source.read_text(encoding="utf-8-sig"))
        return isinstance(data, dict) and isinstance(data.get("features"), list)
    except (ValueError, OSError, ImportError):
        return False
